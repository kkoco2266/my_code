import pandas as pd
import numpy as np
import docx

def database_corrector0(pre_data :str, post_data :str) -> None:
    '''
    This version of the function creates a new file with the processed data
    pre_data is a string with the file location and
    post_data is a string with the name you want the processed data to have
    '''
    data = pd.read_csv(pre_data, header=None).dropna(axis=1)
    proc = pd.concat([pd.DataFrame(data[0].values +' '+ data[1].values), data[[2]]], axis=1).set_index(2)
    proc.rename(columns={0 : 'Student Name'}, inplace=True)
    proc.index.name = 'ID Number'
    proc.index = '0' + proc.index.astype('str')
    proc.to_csv(post_data)

def database_corrector1(pre_data :str) -> pd.DataFrame:
    '''
    This version of the function returns a pandas DataFrame with the processed data
    pre_data is a string with the file location
    '''
    data = pd.read_csv(pre_data, header=None).dropna(axis=1)
    proc = pd.concat([pd.DataFrame(data[0].values +' '+ data[1].values), data[[2]]], axis=1).set_index(2)
    proc.rename(columns={0 : 'Student Name'}, inplace=True)
    proc.index.name = 'ID Number'
    proc.index = '0' + proc.index.astype('str')
    return proc

def proc_obj(obj_filepath :str) -> pd.DataFrame:
    pro_file = pd.read_csv(obj_filepath, index_col='ID Number')
    pro_file = pro_file[pro_file.index.notna()]
    pro_file = pro_file[pro_file['Student Name'].notna()]
    pro_file.index = '0' + pro_file.index.astype('int', copy=False).astype('str', copy=False)
    return pro_file

def proc_subj(pre_data :str, subject_name :str, total_score :int) -> pd.DataFrame:
    '''
    This version of the function returns a pandas DataFrame with the processed data
    pre_data is a string with the file location
    subject_name is the name of the subject and
    total_score is the total score for that test
    '''
    data = pd.read_csv(pre_data, header=None).dropna(axis=1)
    proc = pd.concat([pd.DataFrame(data[0].values +' '+ data[1].values), data.loc[:,2:]], axis=1).set_index(2)
    proc.rename(columns={0 : 'Student Name', 3 : f"{subject_name} subjective score (out of {total_score})"},
                 inplace=True)
    proc.index.name = 'ID Number'
    proc.index = '0' + proc.index.astype('str')
    return proc

def proc_subj1(pre_data :str, subject_name :str, loc_index :int, total_score :int) -> pd.DataFrame:
    '''
    This version of the function returns a pandas DataFrame with the processed data
    pre_data is a string with the file location
    subject_name is the name of the subject 
    loc_index is the index of the subject's column and
    total_score is the total score for that test
    '''
    data = pd.read_csv(pre_data).fillna(0)
    ind_col = data.columns[2]
    old_col_name = data.columns[loc_index]
    proc = data.loc[:,[ind_col,old_col_name]].set_index(ind_col)
    proc.rename(columns={old_col_name : f"{subject_name} subjective score (out of {total_score})"},
                 inplace=True)
    proc.index.name = 'ID Number'
    proc.index = '0' + proc.index.astype('str')
    return proc

def proc_langs0(list_of_lang_files :list) -> pd.DataFrame:
    '''
    This function returns a pandas DataFrame with the processed data
    list_of_lang_files is a list with the language file locations in any order
    '''
    first = True
    for lang in list_of_lang_files:
        pro_file = pd.read_csv(lang, index_col='ID Number')
        pro_file = pro_file[pro_file.index.notna()]
        pro_file = pro_file[pro_file['Student Name'].notna()]
        pro_file.index = '0' + pro_file.index.astype('int', copy=False).astype('str', copy=False)
        if first:
            GL_obj = pro_file
            first = False
        else:
            GL_obj = pd.concat([GL_obj, pro_file], axis=0)
    return GL_obj

def proc_langs1(list_of_lang_files :list) -> pd.DataFrame:
    '''
    This function returns a pandas DataFrame with the processed data
    list_of_lang_files is a list with the language file locations in any order
    '''
    frames = []
    for lang in list_of_lang_files:
        pro_file = pd.read_csv(lang, index_col='ID Number')
        pro_file = pro_file[pro_file.index.notna()]
        pro_file = pro_file[pro_file['Student Name'].notna()]
        pro_file.index = '0' + pro_file.index.astype('int', copy=False).astype('str', copy=False)
        frames.append(pro_file)
    GL_obj = pd.concat(frames, axis=0)
    return GL_obj

def add_obj_n_subj(report :pd.DataFrame, proc_obj :pd.DataFrame, proc_subj :pd.DataFrame) -> tuple:
    words = proc_subj.columns[-1].split(' subjective score (out of ')
    subj_total_score = int(words[1][:-1])
    obj_total_score = len(proc_obj.columns)-6
    subject_name = words[0]
    new_obj_col_name = f"{subject_name} objective Score (out of {obj_total_score})"
    report.loc[:,new_obj_col_name] = proc_obj.loc[:,'Score']
    report.loc[:,proc_subj.columns[-1]] = proc_subj.loc[:,proc_subj.columns[-1]]
    report.fillna(0, inplace=True)
    return obj_total_score, subj_total_score, subject_name

def get_grade(score :float, mark_for_1 :int) -> str:
    if score >= mark_for_1:
        return '1'
    elif score >= mark_for_1 - 5:
        return '2'
    elif score >= mark_for_1 - 10:
        return '3'
    elif score >= mark_for_1 - 15:
        return '4'
    elif score >= mark_for_1 - 20:
        return '5'
    elif score >= mark_for_1 - 25:
        return '6'
    elif score >= mark_for_1 - 30:
        return '7'
    elif score >= mark_for_1 - 35:
        return '8'
    else:
        return '9'

def get_agg(df :pd.DataFrame) -> None:
    sk = [4, 8, 12, 20, 32, 40]
    cores = df.iloc[:,16] + df.iloc[:,24] + df.iloc[:,28]
    for indx in range(len(df)):
        pers = df.iloc[indx].values[sk]
        pers.sort()
        df.loc[df.index[indx],"AGGREGATE"] = pers[:3].sum() + cores.iloc[indx]

class School():
    def __init__(self, students_database :str, list_of_obj_files :list, list_of_subj_files :list,
                 list_of_props :list) -> None:
        '''
        students_database is the file with the students' names and their respective index numbers)

        list_of_obj_files is the list of the objective file locations. The format should be:
        [Career Technology (0),
         Computing (1),
         Creative Arts and Design (2),
         English Language (3),
         French (4),
         Integrated Science (5),
         Mathematics (6),
         RME (7),
         Social Studies (8),
         Ghanaian Language [0 Asante Twi , 1 Akuapem Twi, 2 Ga, 3 Ewe] (9)]
         NB: Ghanaian Language is a nested list with each individual language file location.

        list_of_subj_files is the list of processed subjective score DataFrames.
        They must be already processed with the proc_subj() function.
        The format should be:
        [Career Technology (0),
         Computing (1),
         Creative Arts and Design (2),
         English Language (3),
         French (4),
         Integrated Science (5),
         Mathematics (6),
         RME (7),
         Social Studies (8),
         Ghanaian Language (Asante Twi, Akuapem Twi, Ga, and Ewe combined into a file) (9)]

        list_of_props is a list containing decimals which represent the proportion the objective
        holds in the weighted total
        The format should be:
        [Career Technology (0),
         Computing (1),
         Creative Arts and Design (2),
         English Language (3),
         French (4),
         Integrated Science (5),
         Mathematics (6),
         RME (7),
         Social Studies (8),
         Ghanaian Language (9)]

        '''
        self.database = database_corrector1(students_database)

        self.obj_n_subj_files = [(proc_obj(list_of_obj_files[0]), list_of_subj_files[0]),
                                 (proc_obj(list_of_obj_files[1]), list_of_subj_files[1]),
                                 (proc_obj(list_of_obj_files[2]), list_of_subj_files[2]),
                                 (proc_obj(list_of_obj_files[3]), list_of_subj_files[3]),
                                 (proc_obj(list_of_obj_files[4]), list_of_subj_files[4]),
                                 (proc_obj(list_of_obj_files[5]), list_of_subj_files[5]),
                                 (proc_obj(list_of_obj_files[6]), list_of_subj_files[6]),
                                 (proc_obj(list_of_obj_files[7]), list_of_subj_files[7]),
                                 (proc_obj(list_of_obj_files[8]), list_of_subj_files[8]),
                                 (proc_langs1(list_of_obj_files[9]),list_of_subj_files[9])]
        
        self.props = list_of_props

    def get_report(self, filepath1 :str ='All Subjects Report', filepath2 :str ="All Subjects Compact Report"\
                   , lowest_score_for_1 :int = 75) -> None:
        report = self.database.copy()
        for tupl, obj_prop in zip(self.obj_n_subj_files, self.props) :
            obj_total_score, subj_total_score, subject_name = add_obj_n_subj(report , *tupl)
            obj_scores = report.loc[:,report.columns[-2]]
            subj_scores = report.loc[:,report.columns[-1]]
            report.loc[:,f"{subject_name} total Score (%)"] = \
                ( (obj_scores/obj_total_score)*obj_prop + (subj_scores/subj_total_score)*(1-obj_prop) )*100
            report.loc[:,f"{subject_name} Grade"] = \
                report.loc[:,f"{subject_name} total Score (%)"].apply(lambda x: get_grade(x,lowest_score_for_1))
        report.loc[:,"Total Score (out of 1000)"] = report.iloc[:,np.arange(3,40,4)].sum(axis=1)
        report.loc[:,"Average Score (out of 100)"] = report.loc[:,"Total Score (out of 1000)"]/10
        report.loc[:,"Position"] = \
            report.loc[:,"Total Score (out of 1000)"].rank(method='min',ascending=False, na_option='bottom').astype('int')
        report.iloc[:,1:-1] =report.iloc[:,1:-1].round(decimals=2)
        report.iloc[:,1:] = report.iloc[:,1:].round()
        for col in report.columns[1:]:
            report[col] = report[col].astype(int)
        get_agg(report)
        inds = [0, 3, 4, 7, 8, 11, 12, 15, 16, 19, 20, 23, 24, 27, 28, 31, 32, 35, 36, 39, 40, 41, 44]
        second = report.iloc[:,inds]
        self.com_report = second
        # report.to_csv(filepath1)
        # second.to_csv(filepath2)
        report.to_excel(f"Reports\{filepath1}.xlsx")
        second.to_excel(f"Reports\{filepath2}.xlsx")
    
    def gen_personal_reports(self, sch_name):
        df = self.com_report
        for ind in range(len(df)):
            pers = df.iloc[ind]
            pers.iloc[-1] = pers.iloc[-1].astype(int)
            doc = docx.Document()
            t = doc.add_table(rows=19, cols=3)
            t.cell(0, 1).text = str(sch_name)
            t.cell(2, 0).text = "Student Name :"
            t.cell(2, 1).text = str(pers.iloc[0])
            t.cell(3, 0).text = "Index Number :"
            t.cell(3, 1).text = str(pers.name)
            t.cell(5, 0).text = "Subject Name"
            t.cell(5, 1).text = "Total Score (%)"
            t.cell(5, 2).text = "Grade"
            for ac,dc in zip(np.arange(1,20,2), range(6,16)):
                t.cell(dc, 0).text = str(pers.index[ac].split(" total")[0])
            for ts,rw in zip(np.arange(2,21,2), range(6,16)):
                for j in [1,2]:
                    cell = pers.iloc[ts - j%2]
                    t.cell(rw, j).text = str(cell)
            t.cell(17, 0).text = str("Raw Score :")
            t.cell(17, 1).text = str(pers.iloc[-2])
            t.cell(18, 0).text = str("Aggregate :")
            t.cell(18, 1).text = str(pers.iloc[-1])
            doc.save(f"Reports\personal reports\{pers.iloc[0]}'s Mock 2 Report.docx")
